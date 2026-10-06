import json
import os
import re
import time

from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import requests

from django.utils import timezone
from requests.exceptions import RequestException

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from frostapp.models import Store


ONEC_POSITION_GROUPS = {
    "директор": "Директор",
    "директор магазина": "Директор",

    "администратор": "Администратор",
    "администратор магазина": "Администратор",

    "приемщик": "Приемщик",
    "приемщик товара": "Приемщик",
}

SHARED_ROLE_MARKERS = {
    "administrator": {
        "администратор",
        "админ",
    },
    "receiver": {
        "приемщик",
        "приемка",
    },
}

BITRIX_PAGE_TIMEOUT = int(
    os.getenv("DIRECTORS_REPORT_BITRIX_TIMEOUT", "60")
)

BITRIX_MAX_PAGES = int(
    os.getenv("DIRECTORS_REPORT_BITRIX_MAX_PAGES", "1000")
)

PERSON_EXACT_SCORE = 85
PERSON_SIMILAR_SCORE = 68
SHARED_EXACT_SCORE = 85
SHARED_SIMILAR_SCORE = 68


class DirectorsReportError(RuntimeError):
    pass


def _string(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, (list, tuple, set)):
        return " ".join(
            _string(item)
            for item in value
            if _string(item)
        ).strip()

    return str(value).strip()


def _normalize_text(value: Any) -> str:
    text = _string(value).casefold().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def _normalize_inn(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        for item in value:
            result = _normalize_inn(item)
            if result:
                return result
        return ""

    digits = re.sub(r"\D+", "", _string(value))
    return digits if len(digits) in (10, 12) else ""


def _to_int(value: Any) -> int | None:
    try:
        text = _string(value)
        if not text:
            return None
        return int(text)
    except (TypeError, ValueError):
        return None


def _bool_from_bitrix(value: Any) -> bool:
    return _string(value).upper() in {
        "Y",
        "YES",
        "TRUE",
        "1",
    }


def _extract_onec_rows(payload: Any) -> list[dict]:
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]

    if not isinstance(payload, dict):
        raise DirectorsReportError(
            "1С вернула ответ неизвестного формата."
        )

    for key in (
        "result",
        "data",
        "employees",
        "Сотрудники",
    ):
        value = payload.get(key)
        if isinstance(value, list):
            return [
                item
                for item in value
                if isinstance(item, dict)
            ]

    raise DirectorsReportError(
        "В ответе 1С не найден список сотрудников."
    )


def _load_onec_directors(
    *,
    url: str,
    username: str,
    password: str,
    timeout: int,
) -> tuple[list[dict], int]:
    if not url:
        raise DirectorsReportError(
            "Не настроен ONEC_WORKING_EMPLOYEES_URL."
        )

    auth = None
    if username:
        auth = (username, password)

    try:
        response = requests.get(
            url,
            auth=auth,
            timeout=(10, max(30, timeout)),
        )
        response.raise_for_status()
    except RequestException as exc:
        raise DirectorsReportError(
            f"Не удалось получить сотрудников из 1С: {exc}"
        ) from exc

    try:
        payload = response.json()
    except ValueError as exc:
        raise DirectorsReportError(
            "1С вернула ответ, который не является JSON."
        ) from exc

    all_rows = _extract_onec_rows(payload)
    result = []
    seen = set()

    for item in all_rows:

        state = _normalize_text(item.get("Состояние"))

        # Если поле заполнено, берём только работающих.
        if state and state not in {
            "работа",
            "работает",
        }:
            continue
        position_raw = _string(item.get("Должность"))
        position_normalized = _normalize_text(position_raw)
        
        role_group = ONEC_POSITION_GROUPS.get(
            position_normalized
        )
        
        if not role_group:
            continue

        surname = _string(item.get("Фамилия"))
        first_name = _string(item.get("Имя"))
        middle_name = _string(item.get("Отчество"))

        fio = " ".join(
            part
            for part in (
                surname,
                first_name,
                middle_name,
            )
            if part
        ).strip()

        if not fio:
            fio = _string(item.get("ФИО"))

        inn = _normalize_inn(item.get("ИНН"))
        smstore = _to_int(item.get("ИдМагазина"))

        dedup_key = (
            inn or _normalize_text(fio),
            smstore,
            position_normalized,
        )

        if dedup_key in seen:
            continue

        seen.add(dedup_key)

        result.append({
            "inn": inn,
            "surname": surname,
            "first_name": first_name,
            "middle_name": middle_name,
            "fio": fio,
            "position": position_raw,
            "role_group": role_group,
            "smstore": smstore,
            "email": _string(item.get("Почта")),
            "phone": _string(item.get("НомерТелефона")),
            "department": _string(item.get("Подразделение")),
            "department_guid": _string(
                item.get("ПодразделениеGuid")
            ),
            "organization": _string(
                item.get("ОрганизацияНаименование")
            ),
        })

    return result, len(all_rows)


def _bitrix_request(
    session: requests.Session,
    url: str,
    data: dict,
) -> dict:
    last_error = None

    for attempt in range(1, 5):
        try:
            response = session.post(
                url,
                data=data,
                timeout=(10, BITRIX_PAGE_TIMEOUT),
            )

            body = response.json()

            error_code = _string(body.get("error")).upper()

            if (
                response.status_code == 429
                or error_code in {
                    "QUERY_LIMIT_EXCEEDED",
                    "OPERATION_TIME_LIMIT",
                }
            ):
                retry_after = response.headers.get(
                    "Retry-After",
                    "",
                )
                try:
                    wait_seconds = max(
                        1.0,
                        float(retry_after),
                    )
                except (TypeError, ValueError):
                    wait_seconds = float(attempt)

                time.sleep(min(wait_seconds, 5.0))
                continue

            response.raise_for_status()

            if body.get("error"):
                description = (
                    body.get("error_description")
                    or body.get("error")
                )
                raise DirectorsReportError(
                    f"Ошибка Bitrix24: {description}"
                )

            return body

        except DirectorsReportError:
            raise

        except (
            RequestException,
            ValueError,
        ) as exc:
            last_error = exc

            if attempt < 4:
                time.sleep(float(attempt))
                continue

    raise DirectorsReportError(
        f"Не удалось выполнить запрос к Bitrix24: "
        f"{last_error}"
    )


def _load_all_bitrix_users(
    *,
    url: str,
) -> list[dict]:
    if not url:
        raise DirectorsReportError(
            "Не настроен BITRIX_USER_GET_URL."
        )

    session = requests.Session()
    users = []
    start = 0
    visited_starts = set()

    for _page_number in range(BITRIX_MAX_PAGES):
        if start in visited_starts:
            raise DirectorsReportError(
                "Bitrix24 вернул циклическую пагинацию."
            )

        visited_starts.add(start)

        body = _bitrix_request(
            session,
            url,
            {"start": start},
        )

        page = body.get("result", [])

        if not isinstance(page, list):
            raise DirectorsReportError(
                "Bitrix24 вернул некорректный список "
                "пользователей."
            )

        users.extend(
            user
            for user in page
            if isinstance(user, dict)
        )

        next_start = body.get("next")

        if next_start in (None, ""):
            break

        try:
            start = int(next_start)
        except (TypeError, ValueError) as exc:
            raise DirectorsReportError(
                "Bitrix24 вернул некорректное значение next."
            ) from exc
    else:
        raise DirectorsReportError(
            "Превышено допустимое количество страниц "
            "Bitrix24."
        )

    return users


def _bitrix_display_name(user: dict) -> str:
    fio = " ".join(
        value
        for value in (
            _string(user.get("LAST_NAME")),
            _string(user.get("NAME")),
            _string(user.get("SECOND_NAME")),
        )
        if value
    ).strip()

    if fio:
        return fio

    return (
        _string(user.get("EMAIL"))
        or f"Пользователь ID {user.get('ID', '?')}"
    )


def _prepare_bitrix_users(
    users: list[dict],
    inn_field: str,
) -> list[dict]:
    prepared = []

    for user in users:
        is_active = _bool_from_bitrix(
            user.get("ACTIVE", "Y")
        )
    
        # Неактивных и уволенных пользователей Bitrix24
        # вообще не рассматриваем как совпадения.
        if not is_active:
            continue
    
        name_text = _normalize_text(
            " ".join([
                _string(user.get("NAME")),
                _string(user.get("LAST_NAME")),
                _string(user.get("SECOND_NAME")),
            ])
        )

        search_text = _normalize_text(
            " ".join([
                name_text,
                _string(user.get("WORK_POSITION")),
                _string(user.get("EMAIL")),
            ])
        )

        prepared.append({
            "raw": user,
            "id": _string(user.get("ID")),
            "display": _bitrix_display_name(user),
            "active": True,
            "email": _string(user.get("EMAIL")),
            "work_position": _string(
                user.get("WORK_POSITION")
            ),
            "inn": _normalize_inn(user.get(inn_field)),
            "name_text": name_text,
            "search_text": search_text,
            "name_tokens": set(name_text.split()),
            "search_tokens": set(search_text.split()),
        })

    return prepared


def _best_token_ratio(
    expected: str,
    candidate_tokens: set[str],
) -> float:
    expected = _normalize_text(expected)

    if not expected or not candidate_tokens:
        return 0.0

    return max(
        SequenceMatcher(
            None,
            expected,
            token,
        ).ratio()
        for token in candidate_tokens
    )


def _personal_match_score(
    employee: dict,
    bitrix_user: dict,
) -> tuple[int, str]:
    employee_inn = employee["inn"]
    bitrix_inn = bitrix_user["inn"]

    surname_tokens = set(
        _normalize_text(employee["surname"]).split()
    )
    first_name_tokens = set(
        _normalize_text(employee["first_name"]).split()
    )
    middle_name_tokens = set(
        _normalize_text(employee["middle_name"]).split()
    )

    user_tokens = bitrix_user["name_tokens"]

    surname_exact = (
        bool(surname_tokens)
        and surname_tokens.issubset(user_tokens)
    )
    first_name_exact = (
        bool(first_name_tokens)
        and first_name_tokens.issubset(user_tokens)
    )
    middle_name_exact = (
        not middle_name_tokens
        or middle_name_tokens.issubset(user_tokens)
    )

    inn_exact = (
        bool(employee_inn)
        and bool(bitrix_inn)
        and employee_inn == bitrix_inn
    )

    if inn_exact and surname_exact and first_name_exact:
        return 100, "ИНН и ФИО"

    if inn_exact:
        return 96, "ИНН"

    if surname_exact and first_name_exact:
        if (
            employee_inn
            and bitrix_inn
            and employee_inn != bitrix_inn
        ):
            return 78, "ФИО, но ИНН отличается"

        if middle_name_exact:
            return 92, "Полное ФИО"

        return 88, "Фамилия и имя"

    surname_ratio = _best_token_ratio(
        employee["surname"],
        user_tokens,
    )
    first_name_ratio = _best_token_ratio(
        employee["first_name"],
        user_tokens,
    )

    if surname_ratio >= 0.88 and first_name_ratio >= 0.88:
        score = round(
            68
            + min(surname_ratio, first_name_ratio) * 10
        )
        return score, "Похожее ФИО"

    return 0, ""


def _public_match(
    bitrix_user: dict,
    score: int,
    matched_by: str,
) -> dict:
    return {
        "id": bitrix_user["id"],
        "display": bitrix_user["display"],
        "active": bitrix_user["active"],
        "email": bitrix_user["email"],
        "work_position": bitrix_user["work_position"],
        "inn": bitrix_user["inn"],
        "score": score,
        "matched_by": matched_by,
    }


def _match_person(
    employee: dict,
    bitrix_users: list[dict],
) -> tuple[list[dict], list[dict]]:
    exact = []
    similar = []

    for bitrix_user in bitrix_users:
        score, matched_by = _personal_match_score(
            employee,
            bitrix_user,
        )

        if not score:
            continue

        match = _public_match(
            bitrix_user,
            score,
            matched_by,
        )

        if score >= PERSON_EXACT_SCORE:
            exact.append(match)
        elif score >= PERSON_SIMILAR_SCORE:
            similar.append(match)

    sort_key = lambda item: (
        -item["score"],
        not item["active"],
        item["display"].casefold(),
        item["id"],
    )

    exact.sort(key=sort_key)
    similar.sort(key=sort_key)

    return exact, similar[:10]


def _load_store_alias_overrides() -> dict[str, list[str]]:
    raw = os.getenv(
        "DIRECTORS_STORE_ALIASES_JSON",
        "",
    ).strip()

    if not raw:
        return {}

    try:
        payload = json.loads(raw)
    except ValueError as exc:
        raise DirectorsReportError(
            "DIRECTORS_STORE_ALIASES_JSON содержит "
            "некорректный JSON."
        ) from exc

    if not isinstance(payload, dict):
        raise DirectorsReportError(
            "DIRECTORS_STORE_ALIASES_JSON должен быть "
            "JSON-объектом."
        )

    result = {}

    for key, value in payload.items():
        if isinstance(value, str):
            aliases = [value]
        elif isinstance(value, list):
            aliases = [
                _string(item)
                for item in value
                if _string(item)
            ]
        else:
            continue

        if aliases:
            result[str(key)] = aliases

    return result


LOCATION_STOP_WORDS = {
    "магазин",
    "дискаунтер",
    "николаевский",
    "улица",
    "ул",
    "дом",
    "проспект",
    "пр",
    "переулок",
    "район",
    "город",
    "село",
    "поселок",
    "пос",
}


def _tokenize_account_name(value: Any) -> list[str]:
    """
    Разделяет буквы и цифры.

    Например:
      Дискаунтер15 -> ["дискаунтер", "15"]
      Шилка-2      -> ["шилка", "2"]
    """
    text = _string(value).casefold().replace("ё", "е")

    return re.findall(
        r"[a-zа-я]+|\d+",
        text,
        flags=re.IGNORECASE,
    )


def _location_words(*values: Any) -> list[str]:
    result = set()

    for value in values:
        for token in _normalize_text(value).split():
            if (
                token.isalpha()
                and len(token) >= 4
                and token not in LOCATION_STOP_WORDS
            ):
                result.add(token)

    return sorted(result)




ADDRESS_STOP_WORDS = {
    "магазин",
    "дискаунтер",
    "николаевский",
    "улица",
    "ул",
    "дом",
    "проспект",
    "пр",
    "переулок",
    "район",
    "город",
    "село",
    "поселок",
    "корпус",
    "строение",
    "улан",
    "удэ",
    "чита",
    "россия",
    "республика",
    "бурятия",
    "забайкальский",
    "край",
}

SMALL_ADDRESS_NUMBER_RE = re.compile(
    r"(?<!\d)"
    r"(\d{1,3}[a-zа-я]?"
    r"(?:[/\-]\d{1,3}[a-zа-я]?)?)"
    r"(?!\d)",
    flags=re.IGNORECASE,
)


def _address_words(value: Any) -> set[str]:
    result = set()

    for token in _normalize_text(value).split():
        if (
            token.isalpha()
            and len(token) >= 4
            and token not in ADDRESS_STOP_WORDS
        ):
            result.add(token)

    return result


def _address_numbers(value: Any) -> set[str]:
    text = (
        _string(value)
        .casefold()
        .replace("ё", "е")
        .replace(" ", "")
    )

    return {
        match.casefold()
        for match in SMALL_ADDRESS_NUMBER_RE.findall(
            text
        )
    }


def _extract_store_house(address: Any) -> str:
    text = (
        _string(address)
        .casefold()
        .replace("ё", "е")
    )

    # Сначала ищем явное обозначение дома.
    explicit_match = re.search(
        r"(?:\bдом\b|\bд\.)\s*"
        r"(\d{1,3}[a-zа-я]?"
        r"(?:[/\-]\d{1,3}[a-zа-я]?)?)",
        text,
        flags=re.IGNORECASE,
    )

    if explicit_match:
        return explicit_match.group(1).replace(
            " ",
            "",
        )

    numbers = list(
        SMALL_ADDRESS_NUMBER_RE.findall(text)
    )

    if numbers:
        return numbers[-1].casefold()

    return ""


def _address_match_level(
    identity: dict,
    account_name: str,
) -> tuple[str, list[str]]:
    store_words = set(
        identity.get("address_words", [])
    )
    store_house = identity.get("address_house", "")

    account_words = _address_words(account_name)
    account_numbers = _address_numbers(account_name)

    common_words = sorted(
        store_words & account_words
    )

    street_matches = bool(common_words)
    house_matches = bool(
        store_house
        and store_house in account_numbers
    )

    if street_matches and house_matches:
        return "exact", common_words

    if street_matches or house_matches:
        return "partial", common_words

    return "none", []

def _build_store_identity(
    store: dict | None,
    smstore: int | None,
    overrides: dict[str, list[str]],
) -> dict:
    identity = {
        "label": "",
        "brand": "",
        "number": None,
        "aliases": [],
        "address_words": [],
        "address_house": "",
        "recognized": False,
    }

    if not store:
        return identity

    raw_name = _string(store.get("name"))
    raw_address = _string(store.get("address"))

    normalized_name = (
        raw_name.casefold().replace("ё", "е")
    )

    remaining_name = normalized_name

    brand_match = re.search(
        r"(николаевский|дискаунтер)",
        normalized_name,
        flags=re.IGNORECASE,
    )

    if brand_match:
        identity["brand"] = brand_match.group(1)

        remaining_name = normalized_name[
            brand_match.end():
        ].strip(" \t\r\n-–—№#")

        numeric_match = re.match(
            r"(\d+)",
            remaining_name,
        )

        if numeric_match:
            raw_number = numeric_match.group(1)

            identity["number"] = int(raw_number)
            identity["label"] = raw_number
            identity["recognized"] = True

            remaining_name = remaining_name[
                numeric_match.end():
            ].strip(" \t\r\n-–—№#")

        else:
            location_match = re.match(
                r"([a-zа-я]+(?:\s*[-–—]\s*\d+)?)",
                remaining_name,
                flags=re.IGNORECASE,
            )

            if location_match:
                alias = location_match.group(1).strip()

                identity["aliases"].append(alias)
                identity["label"] = alias
                identity["recognized"] = True

                remaining_name = remaining_name[
                    location_match.end():
                ].strip()

    address_source = raw_address or remaining_name

    identity["address_words"] = sorted(
        _address_words(address_source)
    )
    identity["address_house"] = (
        _extract_store_house(address_source)
    )

    for alias in overrides.get(str(smstore), []):
        alias = _string(alias)

        if alias and alias not in identity["aliases"]:
            identity["aliases"].append(alias)

        if alias and not identity["label"]:
            identity["label"] = alias

        if alias:
            identity["recognized"] = True

    return identity


def _shared_role_positions(
    tokens: list[str],
    shared_role: str,
) -> list[int]:
    markers = SHARED_ROLE_MARKERS.get(
        shared_role,
        set(),
    )

    return [
        index
        for index, token in enumerate(tokens)
        if token in markers
    ]


def _nearest_numbers_to_roles(
    tokens: list[str],
    role_positions: list[int],
) -> set[int]:
    """
    Берёт ближайшее небольшое число около названия роли.

    Это позволяет отличить:
      Гагарина 15, 441210 30 Администратор
    где номер магазина — 30, а 15 — номер дома.
    """
    result = set()

    for role_position in role_positions:
        for distance in range(1, 5):
            found_at_distance = []

            left_index = role_position - distance
            right_index = role_position + distance

            for index in (left_index, right_index):
                if index < 0 or index >= len(tokens):
                    continue

                token = tokens[index]

                if not token.isdigit():
                    continue

                number = int(token)

                # Номера магазинов небольшие.
                # Телефоны и другие длинные числа исключаем.
                if 1 <= number <= 999:
                    found_at_distance.append(number)

            if found_at_distance:
                result.update(found_at_distance)
                break

    return result


def _alias_tokens(alias: str) -> set[str]:
    return set(_tokenize_account_name(alias))


def _best_alias_ratio(
    alias: str,
    bitrix_user: dict,
) -> float:
    """
    Сравнивает название магазина с отдельными фрагментами
    имени пользователя Bitrix24.

    Используется для неточного поиска магазинов:
    Каштак, Шилка-1, Шилка-2 и аналогичных.
    """
    normalized_alias = _normalize_text(alias)
    alias_tokens = normalized_alias.split()

    candidate_tokens = (
        bitrix_user.get("name_text", "")
        .split()
    )

    if not alias_tokens or not candidate_tokens:
        return 0.0

    window_size = len(alias_tokens)
    windows = []

    for index in range(len(candidate_tokens)):
        window = " ".join(
            candidate_tokens[
                index:index + window_size
            ]
        ).strip()

        if window:
            windows.append(window)

    if not windows:
        return 0.0

    return max(
        SequenceMatcher(
            None,
            normalized_alias,
            window,
        ).ratio()
        for window in windows
    )





def _shared_match_score(
    identity: dict,
    bitrix_user: dict,
    shared_role: str,
) -> tuple[int, str]:
    account_name = bitrix_user["name_text"]
    tokens = _tokenize_account_name(account_name)
    token_set = set(tokens)

    role_positions = _shared_role_positions(
        tokens,
        shared_role,
    )

    if not role_positions:
        return 0, ""

    role_title = (
        "Администратор"
        if shared_role == "administrator"
        else "Приемщик/Приемка"
    )

    address_level, common_address_words = (
        _address_match_level(
            identity,
            account_name,
        )
    )

    address_description = ", ".join(
        common_address_words
    )

    alias_matched = False

    for alias in identity.get("aliases", []):
        expected_tokens = _alias_tokens(alias)

        if (
            expected_tokens
            and expected_tokens.issubset(token_set)
        ):
            alias_matched = True
            break

    store_number = identity.get("number")

    if store_number is not None:
        contextual_numbers = _nearest_numbers_to_roles(
            tokens,
            role_positions,
        )

        # Номер возле роли должен совпадать с номером
        # магазина, а не с номером дома.
        if store_number not in contextual_numbers:
            return 0, ""

        if address_level == "exact":
            return (
                100,
                f"{role_title}: номер магазина, "
                f"улица и дом совпали"
                + (
                    f" ({address_description})"
                    if address_description
                    else ""
                ),
            )

        if address_level == "partial":
            return (
                78,
                f"{role_title}: номер магазина совпал, "
                f"адрес совпал только частично",
            )

        return (
            72,
            f"{role_title}: номер магазина совпал, "
            f"но адрес не подтверждён",
        )

    if alias_matched:
        # Для Шилки, Каштака и других именованных
        # магазинов также проверяем адрес.
        if address_level == "exact":
            return (
                100,
                f"{role_title}: название магазина, "
                f"улица и дом совпали",
            )

        if not identity.get("address_words"):
            return (
                96,
                f"{role_title}: совпало уникальное "
                f"название магазина",
            )

        return (
            78,
            f"{role_title}: название магазина совпало, "
            f"но адрес требует проверки",
        )

    for alias in identity.get("aliases", []):
        ratio = _best_alias_ratio(
            alias,
            bitrix_user,
        )

        if ratio >= 0.84:
            return (
                70,
                f"{role_title}: похожее название "
                f"магазина",
            )

    return 0, ""


def _match_shared_account(
    identity: dict,
    bitrix_users: list[dict],
    shared_role: str,
) -> tuple[list[dict], list[dict]]:
    if not identity.get("recognized"):
        return [], []

    exact = []
    similar = []

    for bitrix_user in bitrix_users:
        score, matched_by = _shared_match_score(
            identity,
            bitrix_user,
            shared_role,
        )

        if not score:
            continue

        match = _public_match(
            bitrix_user,
            score,
            matched_by,
        )

        if score >= SHARED_EXACT_SCORE:
            exact.append(match)
        elif score >= SHARED_SIMILAR_SCORE:
            similar.append(match)

    sort_key = lambda item: (
        -item["score"],
        not item["active"],
        item["display"].casefold(),
        item["id"],
    )

    exact.sort(key=sort_key)
    similar.sort(key=sort_key)

    return exact, similar[:10]


def _compact_account_status(
    exact: list[dict],
    similar: list[dict],
) -> str:
    if len(exact) == 1:
        return "Найдено"

    if len(exact) > 1:
        return "Несколько совпадений"

    if similar:
        return "Требует проверки"

    return "Не найдено"


def _personal_status(
    exact: list[dict],
    similar: list[dict],
) -> str:
    if len(exact) == 1:
        return "Найдено"

    if len(exact) > 1:
        return "Несколько совпадений"

    if similar:
        return "Требует проверки"

    return "Не найдено"


def _shared_status(
    exact: list[dict],
    similar: list[dict],
) -> str:
    if exact:
        return f"Найдено: {len(exact)}"

    if similar:
        return "Требует проверки"

    return "Не найдено"


def build_directors_report(
    *,
    onec_url: str,
    onec_username: str,
    onec_password: str,
    onec_timeout: int,
    bitrix_user_get_url: str,
    bitrix_inn_field: str,
) -> dict:
    employees, onec_total = _load_onec_directors(
        url=onec_url,
        username=onec_username,
        password=onec_password,
        timeout=onec_timeout,
    )

    bitrix_raw_users = _load_all_bitrix_users(
        url=bitrix_user_get_url,
    )

    # Здесь уже остаются только активные пользователи.
    bitrix_users = _prepare_bitrix_users(
        bitrix_raw_users,
        bitrix_inn_field,
    )

    store_ids = {
        employee["smstore"]
        for employee in employees
        if employee["smstore"] is not None
    }

    stores_by_smstore = defaultdict(list)

    for store_model in Store.objects.filter(
        smstore__in=store_ids
    ).only(
        "id",
        "smstore",
        "name",
        "region",
        "address",
        "close_date",
    ):
        stores_by_smstore[
            store_model.smstore
        ].append(store_model)

    prepared_stores = {}

    for smstore, store_models in stores_by_smstore.items():
        # Сначала действующий магазин, затем самая
        # новая запись.
        store_models.sort(
            key=lambda item: (
                item.close_date is not None,
                -(item.id or 0),
            )
        )

        selected = store_models[0]

        prepared_stores[smstore] = {
            "id": selected.id,
            "smstore": selected.smstore,
            "name": _string(selected.name),
            "region": _string(selected.region),
            "address": _string(selected.address),
            "close_date": (
                selected.close_date.isoformat()
                if selected.close_date
                else ""
            ),
            "duplicate_count": len(store_models),
        }

    aliases = _load_store_alias_overrides()
    shared_cache = {}
    rows = []

    status_totals = {
        "found": 0,
        "ambiguous": 0,
        "possible": 0,
        "missing": 0,
    }

    for employee in employees:
        smstore = employee["smstore"]
        store = prepared_stores.get(smstore)

        store_name = (
            (store or {}).get("name")
            or f"smstore={smstore}"
        )

        comments = []

        if smstore is None:
            comments.append(
                "В 1С не заполнен ИдМагазина."
            )
        elif not store:
            comments.append(
                f"В stores не найден smstore={smstore}."
            )
        else:
            if store["duplicate_count"] > 1:
                comments.append(
                    "В stores несколько записей "
                    "с одинаковым smstore."
                )

            if store["close_date"]:
                comments.append(
                    "У магазина указана дата закрытия "
                    f"{store['close_date']}."
                )

        role_group = employee["role_group"]

        if role_group == "Директор":
            account_type = "Личная"

            exact, similar = _match_person(
                employee,
                bitrix_users,
            )

        else:
            if role_group == "Администратор":
                shared_role = "administrator"
                account_type = "Общая администратора"
            else:
                shared_role = "receiver"
                account_type = "Общая приемки"

            cache_key = (
                smstore,
                shared_role,
            )

            if cache_key not in shared_cache:
                identity = _build_store_identity(
                    store,
                    smstore,
                    aliases,
                )

                exact, similar = (
                    _match_shared_account(
                        identity,
                        bitrix_users,
                        shared_role,
                    )
                )

                shared_cache[cache_key] = {
                    "exact": exact,
                    "similar": similar,
                    "identity": identity,
                }

            cached = shared_cache[cache_key]

            exact = cached["exact"]
            similar = cached["similar"]
            identity = cached["identity"]

            if store and not identity["recognized"]:
                comments.append(
                    "Не удалось определить номер или "
                    "название магазина."
                )

        status = _compact_account_status(
            exact,
            similar,
        )

        if status == "Найдено":
            status_totals["found"] += 1

            comments.append(
                exact[0]["matched_by"]
            )

        elif status == "Несколько совпадений":
            status_totals["ambiguous"] += 1

            comments.append(
                "Найдено несколько точных записей. "
                "Нужно проверить дубли."
            )

        elif status == "Требует проверки":
            status_totals["possible"] += 1

            comments.append(
                "Есть похожая запись, но совпадение "
                "не подтверждено полностью."
            )

        else:
            status_totals["missing"] += 1

            comments.append(
                "Подходящая активная учётная запись "
                "не найдена."
            )

        rows.append({
            "store": store,
            "smstore": smstore,
            "store_name": store_name,
            "position": employee["position"],
            "role_group": role_group,
            "fio": employee["fio"],
            "inn": employee["inn"],
            "account_type": account_type,
            "status": status,
            "exact": exact,
            "similar": similar,
            "comment": " ".join(comments),
        })

    role_order = {
        "Директор": 1,
        "Администратор": 2,
        "Приемщик": 3,
    }

    rows.sort(
        key=lambda row: (
            _string(
                (row.get("store") or {}).get("region")
            ).casefold(),
            row["store_name"].casefold(),
            role_order.get(row["role_group"], 9),
            row["fio"].casefold(),
        )
    )

    summary = {
        "onec_total_employees": onec_total,
        "target_employees": len(rows),
        "stores_1c": len({
            row["smstore"]
            for row in rows
        }),
        "directors_1c": sum(
            row["role_group"] == "Директор"
            for row in rows
        ),
        "administrators_1c": sum(
            row["role_group"] == "Администратор"
            for row in rows
        ),
        "receivers_1c": sum(
            row["role_group"] == "Приемщик"
            for row in rows
        ),
        "bitrix_users": len(bitrix_users),
        "found": status_totals["found"],
        "ambiguous": status_totals["ambiguous"],
        "possible": status_totals["possible"],
        "missing": status_totals["missing"],
        "issues": (
            status_totals["ambiguous"]
            + status_totals["possible"]
            + status_totals["missing"]
        ),
    }

    return {
        "generated_at": timezone.localtime().isoformat(),
        "summary": summary,
        "rows": rows,
    }


def _excel_safe(value: Any) -> Any:
    if value is None:
        return ""

    if not isinstance(value, str):
        return value

    if value.startswith(("=", "+", "-", "@")):
        return "'" + value

    return value


def _matches_to_text(
    exact: list[dict],
    similar: list[dict],
) -> str:
    parts = []

    for match in exact:
        active = "активен" if match["active"] else "неактивен"
        parts.append(
            f"ТОЧНО: {match['display']} "
            f"[ID {match['id']}; {active}; "
            f"{match['matched_by']}]"
        )

    for match in similar:
        active = "активен" if match["active"] else "неактивен"
        parts.append(
            f"ПОХОЖЕ: {match['display']} "
            f"[ID {match['id']}; {active}; "
            f"{match['matched_by']}]"
        )

    return "\n".join(parts)


def _style_worksheet(
    worksheet,
    *,
    freeze: str = "A2",
):
    worksheet.freeze_panes = freeze

    if worksheet.max_row >= 1:
        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    thin = Side(
        style="thin",
        color="D9E2F3",
    )

    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = Border(
                left=thin,
                right=thin,
                top=thin,
                bottom=thin,
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):
        max_length = 0

        for row_index in range(
            1,
            min(worksheet.max_row, 300) + 1,
        ):
            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is not None:
                max_length = max(
                    max_length,
                    min(len(str(value)), 60),
                )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = max(12, min(max_length + 2, 60))

    worksheet.auto_filter.ref = worksheet.dimensions


def _employee_names_text(
    store_result: dict,
    role_group: str,
) -> str:
    names = store_result.get(
        "names_by_role",
        {},
    ).get(
        role_group,
        [],
    )

    return "\n".join(names)


def write_directors_report_xlsx(
    report: dict,
    target_path: Path,
):
    target_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Отчёт"

    worksheet.append([
        "Магазин",
        "Должность",
        "ФИО",
        "Учётная запись Bitrix24",
        "Результат / уточнение",
    ])

    for row in report["rows"]:
        store = row.get("store") or {}

        store_parts = [
            row["store_name"],
        ]

        if row.get("smstore") is not None:
            store_parts.append(
                f"smstore={row['smstore']}"
            )

        if store.get("region"):
            store_parts.append(store["region"])

        store_text = "\n".join(store_parts)

        matches = _matches_to_text(
            row["exact"],
            row["similar"],
        )

        account_text = (
            f"{row['account_type']}\n{matches}"
            if matches
            else (
                f"{row['account_type']}\n"
                "Учётная запись не найдена"
            )
        )

        result_text = (
            f"{row['status']}. {row['comment']}"
        ).strip()

        worksheet.append([
            _excel_safe(store_text),
            _excel_safe(row["position"]),
            _excel_safe(row["fio"]),
            _excel_safe(account_text),
            _excel_safe(result_text),
        ])

    _style_worksheet(worksheet)

    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 26
    worksheet.column_dimensions["C"].width = 38
    worksheet.column_dimensions["D"].width = 65
    worksheet.column_dimensions["E"].width = 65

    workbook.save(target_path)
